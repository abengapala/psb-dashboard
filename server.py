#!/usr/bin/env python3
"""
PSB AUTO CURING - dashboard server
==================================
Serves the React TV dashboard (dist/) and gives the admin page an upload endpoint.

  TV screen   : http://<this-PC-IP>:8080/
  Admin page  : http://<this-PC-IP>:8080/#admin   (password protected)

What it does
  * Serves dist/ as plain static files (TV page polls data.json every 15 s).
  * POST /api/login   -> checks the admin password.
  * POST /upload      -> admin uploads the Excel file; it is saved to uploads/ and
                         data.json is rebuilt immediately with update_dashboard.py.
  * Also watches the team's Excel file (--xlsx) and rebuilds whenever it is saved.

It does NOT modify update_dashboard.py; it just imports and calls its build().

Usage
  python server.py
  python server.py --password "my-secret" --port 8080
  python server.py --xlsx "C:/path/PSB AUTO CURING MONITORING 2026.xlsx"

Needs: pip install openpyxl     (and run "npm run build" once to create dist/)
"""
import argparse
import datetime as dt
import functools
import http.server
import importlib.util
import json
import os
import shutil
import socketserver
import sys
import tempfile
import threading
import time
from pathlib import Path
from urllib.parse import unquote

HERE = Path(__file__).resolve().parent
DOWNLOADS = Path.home() / "Downloads"
MAX_UPLOAD = 50 * 1024 * 1024  # 50 MB

try:
    sys.stdout.reconfigure(errors="replace")
except Exception:
    pass


def log(msg):
    print(f"[{dt.datetime.now():%H:%M:%S}] {msg}", flush=True)


def load_generator(path: Path):
    if not path.exists():
        sys.exit(f"update_dashboard.py not found at {path}  (use --script)")
    spec = importlib.util.spec_from_file_location("update_dashboard", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


class App:
    def __init__(self, args):
        self.args = args
        self.dist = args.dist.resolve()
        self.gen = load_generator(args.script.resolve())
        self.gen.load_data = self._load_data  # robust reader (see below)
        self.lock = threading.Lock()
        self.uploads = HERE / "uploads"
        self.uploads.mkdir(exist_ok=True)

    def _load_data(self, xlsx: Path):
        """Same as update_dashboard.load_data, but ignores the sheet-size tag.
        Some saved copies report every sheet as 1 cell (A1), which made the
        original reader see an empty PTP sheet."""
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            copy = Path(tmp) / "snapshot.xlsx"
            shutil.copy2(xlsx, copy)
            wb = load_workbook(copy, data_only=True, read_only=True)
            try:
                out = []
                for name in (self.gen.SHEET_PTP, self.gen.SHEET_PERF):
                    ws = wb[name]
                    ws.reset_dimensions()
                    out.append(list(ws.iter_rows(values_only=True)))
            finally:
                wb.close()
        return out[0], out[1]

    @staticmethod
    def _read_eod(xlsx: Path):
        """EOD REPORT sheet, columns A-H: PTP INVENTORY (A-D) + CONFIRMED PAYMENT (E-H) per bucket."""
        import re
        from openpyxl import load_workbook

        def num(v):
            if isinstance(v, (int, float)):
                return float(v)
            try:
                return float(str(v).replace(",", "").strip())
            except (TypeError, ValueError):
                return 0.0

        with tempfile.TemporaryDirectory() as tmp:
            copy = Path(tmp) / "eod.xlsx"
            shutil.copy2(xlsx, copy)
            wb = load_workbook(copy, data_only=True, read_only=True)
            try:
                name = next((n for n in wb.sheetnames if n.strip().upper().startswith("EOD REPORT")), None)
                if not name:
                    return None
                ws = wb[name]
                ws.reset_dimensions()
                rows = list(ws.iter_rows(max_row=80, max_col=8, values_only=True))
            finally:
                wb.close()

            hi = next((i for i, r in enumerate(rows) if r and str(r[0] or "").strip().upper() == "BUCKET"), None)
            if hi is None:
                return None

            def side(off):
                out = []
                for r in rows[hi + 1:]:
                    r = tuple(r) + (None,) * 8
                    b = str(r[off] or "").strip()
                    if not b:
                        if out:
                            break
                        continue
                    if b.upper().startswith("GRAND TOTAL"):
                        break
                    out.append({"bucket": b.upper(), "count": num(r[off + 1]),
                                "ac": round(num(r[off + 2]), 2), "ob": round(num(r[off + 3]), 2)})
                return out

            ptp_rows = side(0)
            conf_rows = side(4)

            # If CONFIRMED PAYMENT formulas were saved without cached values, evaluate them from PTP MONITORING 2026
            if conf_rows and all(x["count"] == 0 and x["ac"] == 0 and x["ob"] == 0 for x in conf_rows):
                wb_f = load_workbook(copy, data_only=False, read_only=True)
                try:
                    ws_f = wb_f[name]
                    ws_f.reset_dimensions()
                    f_rows = list(ws_f.iter_rows(max_row=hi + 10, max_col=8, values_only=True))
                finally:
                    wb_f.close()
                sample_formula = ""
                for r in f_rows[hi + 1:]:
                    if r and len(r) > 5 and isinstance(r[5], str) and r[5].startswith("="):
                        sample_formula = r[5]
                        break
                if sample_formula:
                    quotes = re.findall(r'"([^"]+)"', sample_formula)
                    month_q = next((q.strip().upper() for q in quotes if "20" in q and "/" not in q), None)
                    date_q = next((q.strip() for q in quotes if "/" in q or "-" in q), None)
                    target_md = None
                    if date_q:
                        for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d"):
                            try:
                                d = dt.datetime.strptime(date_q, fmt)
                                target_md = (d.month, d.day, d.year)
                                break
                            except ValueError:
                                pass
                    wb_p = load_workbook(copy, data_only=True, read_only=True)
                    try:
                        ptp_name = next((n for n in wb_p.sheetnames if "PTP MONITORING" in n.upper()), None)
                        if ptp_name:
                            ws_p = wb_p[ptp_name]
                            ws_p.reset_dimensions()
                            p_rows = list(ws_p.iter_rows(values_only=True))
                            # Columns: M=12 (CONFIRMED AMOUNT), O=14 (CONFIRMED DATE), P=15 (OB), S=18 (MONTH), U=20 (PLACEMENT)
                            for item in conf_rows:
                                b_norm = item["bucket"].replace(" ", "").upper()
                                c_cnt = 0
                                c_ac = 0.0
                                c_ob = 0.0
                                for pr in p_rows[1:]:
                                    if len(pr) <= 20:
                                        continue
                                    u_val = str(pr[20] or "").strip().upper().replace(" ", "")
                                    if u_val != b_norm:
                                        continue
                                    if month_q and str(pr[18] or "").strip().upper() != month_q:
                                        continue
                                    o_val = pr[14]
                                    if target_md:
                                        if isinstance(o_val, (dt.datetime, dt.date)):
                                            if (o_val.month, o_val.day, o_val.year) != target_md:
                                                continue
                                        else:
                                            if str(o_val or "").strip() != date_q:
                                                continue
                                    c_cnt += 1
                                    c_ac += num(pr[12])
                                    c_ob += num(pr[15])
                                item["count"] = float(c_cnt)
                                item["ac"] = round(c_ac, 2)
                                item["ob"] = round(c_ob, 2)
                    finally:
                        wb_p.close()

        def totals(items):
            return {
                "count": int(round(sum(x["count"] for x in items))),
                "ac": round(sum(x["ac"] for x in items), 2),
                "ob": round(sum(x["ob"] for x in items), 2),
            }

        return {
            "sheet": name,
            "ptp": ptp_rows,
            "ptpTotal": totals(ptp_rows),
            "confirmed": conf_rows,
            "confirmedTotal": totals(conf_rows),
        }

    def rebuild(self, xlsx: Path):
        """Run the existing generator; returns (ok, message, warnings)."""
        with self.lock:
            try:
                self.gen.build(xlsx, self.dist / "data.json", self.dist / "admin-report.html", self.args.month)
            except BaseException as e:  # generator uses SystemExit for data errors
                msg = str(e) or e.__class__.__name__
                log(f"build failed: {msg}")
                return False, msg, []
            try:
                data_file = self.dist / "data.json"
                data = json.loads(data_file.read_text(encoding="utf-8"))
                # Match photos in Downloads/photos or dist/photos by full name OR first name
                photo_dirs = [d for d in (Path(self.gen.PHOTOS_DIR), self.dist / "photos") if d.exists()]
                exts = (".png", ".jpg", ".jpeg", ".webp", ".svg")
                changed = False
                for a in data.get("agents", []):
                    full_l = a["name"].lower()
                    first_l = full_l.split()[0] if full_l else ""
                    matched = None
                    for pdir in photo_dirs:
                        for f in pdir.iterdir():
                            if f.suffix.lower() not in exts:
                                continue
                            stem = f.stem.lower()
                            if stem == first_l or stem == full_l:
                                matched = f"photos/{f.name}"
                                if stem == first_l:
                                    break
                        if matched:
                            break
                    if matched and a.get("photo") != matched:
                        a["photo"] = matched
                        changed = True
                eod = self._read_eod(xlsx)
                if eod is not None:
                    data["eod"] = eod
                    changed = True
                out_json = json.dumps(data, indent=2, ensure_ascii=False)
                if changed:
                    data_file.write_text(out_json, encoding="utf-8")
                pub_json = HERE / "public" / "data.json"
                if pub_json.parent.exists():
                    pub_json.write_text(out_json, encoding="utf-8")
                threading.Thread(target=self._push_to_github, args=(xlsx.name,), daemon=True).start()
                k = data["kpi"]
                summary = (f"{data['month']}: {k['manpower']} agents, posted PHP {k['posted']:,.0f}, "
                           f"{k['pctAchieved']:.1f}% of target")
                return True, summary, data.get("warnings", [])
            except Exception as e:
                log(f"post-build warning: {e}")
                return True, "data.json rebuilt", []

    @staticmethod
    def _push_to_github(src_name: str):
        """Auto-commit & push updated public/data.json to GitHub so Vercel updates automatically."""
        import subprocess
        if not (HERE / ".git").exists():
            return
        try:
            subprocess.run(["git", "-C", str(HERE), "add", "public/data.json"], check=False, capture_output=True)
            diff = subprocess.run(["git", "-C", str(HERE), "diff", "--cached", "--quiet"], check=False)
            if diff.returncode != 0:
                subprocess.run(
                    ["git", "-C", str(HERE), "commit", "-m", f"Auto-update data.json from {src_name}"],
                    check=False,
                    capture_output=True,
                )
                res = subprocess.run(["git", "-C", str(HERE), "push", "origin", "main"], check=False, capture_output=True, text=True)
                if res.returncode == 0:
                    log(f"synced data.json to GitHub / Vercel ({src_name})")
                else:
                    log(f"git push warning: {res.stderr.strip()}")
        except Exception as e:
            log(f"git sync error: {e}")


def make_handler(app: App):
    class Handler(http.server.SimpleHTTPRequestHandler):
        def __init__(self, *a, **kw):
            super().__init__(*a, directory=str(app.dist), **kw)

        def log_message(self, *a):
            pass

        # photos/ can live next to update_dashboard.py OR inside dist/photos
        def translate_path(self, path):
            clean = path.split("?", 1)[0].split("#", 1)[0]
            if clean.startswith("/photos/"):
                from urllib.parse import unquote as _unq
                rel = Path(*[p for p in _unq(clean[len("/photos/"):]).split("/") if p not in ("", ".", "..")])
                ext_photo = Path(app.gen.PHOTOS_DIR) / rel
                if ext_photo.exists():
                    return str(ext_photo)
                dist_photo = app.dist / "photos" / rel
                return str(dist_photo)
            return super().translate_path(path)

        def end_headers(self):
            if self.path.split("?", 1)[0].endswith((".json", ".html", "/")):
                self.send_header("Cache-Control", "no-store")
            super().end_headers()

        def _json(self, code, obj):
            body = json.dumps(obj).encode("utf-8")
            self.send_response(code)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def _authed(self, pw=None):
            pw = pw if pw is not None else self.headers.get("X-Admin-Password", "")
            return pw == app.args.password

        def do_POST(self):
            route = self.path.split("?", 1)[0]
            length = int(self.headers.get("Content-Length") or 0)
            if length > MAX_UPLOAD:
                return self._json(413, {"message": "File too large (max 50 MB)."})
            body = self.rfile.read(length) if length else b""

            if route == "/api/login":
                try:
                    pw = json.loads(body or b"{}").get("password", "")
                except Exception:
                    pw = ""
                if self._authed(pw):
                    return self._json(200, {"ok": True})
                time.sleep(1)  # slow down guessing
                return self._json(401, {"ok": False, "message": "Incorrect password."})

            if route == "/upload":
                if not self._authed():
                    return self._json(401, {"message": "Not logged in. Please sign in again."})
                name = Path(unquote(self.headers.get("X-Filename", "upload.xlsx"))).name
                if not name.lower().endswith((".xlsx", ".xlsm")) or not body:
                    return self._json(400, {"message": "Please upload an Excel .xlsx file."})
                stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
                saved = app.uploads / f"{stamp} {name}"
                saved.write_bytes(body)
                log(f"admin uploaded {name} ({len(body)/1024:.0f} KB)")
                ok, msg, warns = app.rebuild(saved)
                if ok:
                    lines = [f"File: {name}", msg, "TV screen will update within 15 seconds."]
                    if warns:
                        lines += ["", "Warnings (fix these in the Excel file):"] + [f"- {w}" for w in warns]
                    return self._json(200, {"message": "Dashboard updated!", "log": lines})
                return self._json(422, {"message": "Could not read this Excel file.", "log": [msg]})

            self._json(404, {"message": "Not found"})

    return Handler


def find_latest_xlsx(xlsx: Path) -> Path | None:
    """Pick the newest matching PSB AUTO CURING MONITORING*.xlsx in the same folder
    (so Windows copy suffixes like '(1)' are picked up automatically)."""
    parent = xlsx.parent
    if not parent.exists():
        return xlsx if xlsx.exists() else None
    candidates = [
        f for f in parent.glob("PSB AUTO CURING MONITORING*.xlsx")
        if f.is_file() and not f.name.startswith("~$")
    ]
    if xlsx.exists() and xlsx not in candidates:
        candidates.append(xlsx)
    if not candidates:
        return None
    return max(candidates, key=lambda f: f.stat().st_mtime)


def watch(app: App, xlsx: Path, interval: int):
    last = None
    while True:
        try:
            target = find_latest_xlsx(xlsx)
            if target and target.exists():
                st = target.stat()
                sig = (str(target), st.st_mtime, st.st_size)
                if sig != last:
                    if last is not None:
                        time.sleep(2)  # let Excel finish writing
                    log(f"Excel changed: {target.name}")
                    app.rebuild(target)
                    last = sig
        except Exception as e:
            log(f"watch error (will retry): {e}")
        time.sleep(interval)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--script", type=Path, default=DOWNLOADS / "update_dashboard.py",
                    help="path to the existing update_dashboard.py")
    ap.add_argument("--xlsx", type=Path, default=DOWNLOADS / "PSB AUTO CURING MONITORING 2026.xlsx",
                    help="team Excel file to watch (optional)")
    ap.add_argument("--dist", type=Path, default=HERE / "dist")
    ap.add_argument("--port", type=int, default=8080)
    ap.add_argument("--month", default=None)
    ap.add_argument("--interval", type=int, default=15)
    ap.add_argument("--password", default=os.environ.get("PSB_ADMIN_PASSWORD", "psb2026"))
    args = ap.parse_args()

    if not (args.dist / "index.html").exists():
        sys.exit(f"{args.dist} has no index.html - run 'npm run build' first.")

    app = App(args)
    threading.Thread(target=watch, args=(app, args.xlsx, args.interval), daemon=True).start()

    socketserver.ThreadingTCPServer.allow_reuse_address = True
    httpd = socketserver.ThreadingTCPServer(("", args.port), make_handler(app))
    log(f"TV dashboard : http://localhost:{args.port}/")
    log(f"Admin page   : http://localhost:{args.port}/#admin")
    log(f"Watching     : {args.xlsx}" + ("" if args.xlsx.exists() else "  (not found yet)"))
    log("Press Ctrl+C to stop.")
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()
